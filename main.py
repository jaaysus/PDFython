import tkinter as tk
from tkinter import filedialog, messagebox
import pdfplumber
from fpdf import FPDF
import openpyxl
import os

class PDFythonApp:
    def __init__(self, root):
        self.root = root
        self.root.title("PDFython - PDF Management")
        self.root.geometry("400x300")

        self.create_widgets()

    def create_widgets(self):
        tk.Button(self.root, text="Convert TXT to PDF", command=self.convert_txt_to_pdf).pack(pady=5)
        tk.Button(self.root, text="Convert Excel to PDF", command=self.convert_excel_to_pdf).pack(pady=5)
        tk.Button(self.root, text="Extract Text from PDF", command=self.extract_text_from_pdf).pack(pady=5)
        tk.Button(self.root, text="Exit", command=self.root.quit).pack(pady=10)

    def convert_txt_to_pdf(self):
        file_path = filedialog.askopenfilename(filetypes=[("Text Files", "*.txt")])
        if not file_path:
            return

        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=10)
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        with open(file_path, "r", encoding="utf-8") as file:
            for line in file:
                pdf.cell(200, 10, line, ln=True)

        save_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")])
        if save_path:
            pdf.output(save_path)
            messagebox.showinfo("Success", "TXT converted to PDF successfully!")

    def convert_excel_to_pdf(self):
        file_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        if not file_path:
            return

        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        for row in ws.iter_rows(values_only=True):
            pdf.cell(200, 10, " | ".join(str(cell) for cell in row), ln=True)

        save_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")])
        if save_path:
            pdf.output(save_path)
            messagebox.showinfo("Success", "Excel converted to PDF successfully!")

    def extract_text_from_pdf(self):
        file_path = filedialog.askopenfilename(filetypes=[("PDF Files", "*.pdf")])
        if not file_path:
            return

        extracted_text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                extracted_text += page.extract_text() + "\n"

        save_path = filedialog.asksaveasfilename(defaultextension=".txt", filetypes=[("Text Files", "*.txt")])
        if save_path:
            with open(save_path, "w", encoding="utf-8") as file:
                file.write(extracted_text)
            messagebox.showinfo("Success", "Text extracted and saved!")

if __name__ == "__main__":
    root = tk.Tk()
    app = PDFythonApp(root)
    root.mainloop()
